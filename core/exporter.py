from reportlab.lib.pagesizes import letter
from docx import Document
from docx.shared import Inches
import os

def create_markdown(text, image_path, output_path):
    """
    Creates a Markdown file with the analysis results.
    """
    try:
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(f"# RxShield Analysis Report\n\n")
            f.write(f"**Image:** {image_path}\n\n")
            f.write(f"## Analysis Results\n\n")
            f.write(text)
        return True
    except Exception as e:
        print(f"Error creating Markdown: {e}")
        return False

def create_word(text, image_path, output_path):
    """
    Creates a Word document with the analysis results.
    """
    try:
        doc = Document()
        doc.add_heading('RxShield Analysis Report', 0)

        if image_path:
            try:
                doc.add_picture(image_path, width=Inches(6))
            except Exception as e:
                doc.add_paragraph(f"[Error adding image: {e}]")

        doc.add_heading('Analysis Results', level=1)
        doc.add_paragraph(text)

        doc.save(output_path)
        return True
    except Exception as e:
        print(f"Error creating Word doc: {e}")
        return False, f"Failed to save Word doc: {e}"
